"""Bulk email verification platform (Flask).

Run:  python app.py   ->  http://127.0.0.1:5000
"""
import csv
import io
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor

from flask import Flask, jsonify, render_template, request, Response

from verifier import Verifier

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB upload cap

MAX_EMAILS = 10000
WORKERS = 40
JOBS = {}  # job_id -> dict


def extract_emails(filename: str, data: bytes):
    """Read emails from a CSV or XLSX upload. Uses the 'email' column if present,
    otherwise the first column."""
    rows = []
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="ignore")
        rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    header = [h.strip().lower() for h in rows[0]]
    col, start = 0, 0
    for i, h in enumerate(header):
        if h in ("email", "e-mail", "email address", "emailaddress", "mail"):
            col, start = i, 1
            break
    else:
        if "@" not in rows[0][0]:  # first row looks like a header
            start = 1
    emails, seen = [], set()
    for r in rows[start:]:
        if col < len(r):
            e = r[col].strip()
            if e and e.lower() not in seen:
                seen.add(e.lower())
                emails.append(e)
    return emails


def run_job(job_id, emails):
    job = JOBS[job_id]
    verifier = Verifier(smtp_enabled=True)
    job["smtp_available"] = verifier.smtp_available
    job["state"] = "running"

    def work(idx_email):
        idx, email = idx_email
        try:
            res = verifier.verify(email)
        except Exception as e:  # never let one address kill the batch
            res = {"email": email, "status": "unknown", "reason": f"Error: {e}",
                   "syntax": None, "domain": None, "mx": None,
                   "catch_all": None, "smtp": None}
        job["results"][idx] = res
        with job["lock"]:
            job["done"] += 1
            job["counts"][res["status"]] += 1

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        list(pool.map(work, enumerate(emails)))
    job["state"] = "finished"
    job["finished_at"] = time.time()


@app.route("/")
def index():
    return render_template("index.html", max_emails=MAX_EMAILS)


@app.route("/api/upload", methods=["POST"])
def upload():
    f = request.files.get("file")
    pasted = request.form.get("emails", "")
    emails = []
    if f and f.filename:
        emails = extract_emails(f.filename, f.read())
    elif pasted.strip():
        emails = [e.strip() for e in pasted.replace(",", "\n").splitlines() if e.strip()]
    if not emails:
        return jsonify(error="No emails found. Upload a CSV/XLSX or paste a list."), 400
    truncated = len(emails) > MAX_EMAILS
    emails = emails[:MAX_EMAILS]
    job_id = uuid.uuid4().hex[:12]
    JOBS[job_id] = {
        "state": "queued", "total": len(emails), "done": 0,
        "results": [None] * len(emails), "lock": threading.Lock(),
        "counts": {"valid": 0, "invalid": 0, "risky": 0, "unknown": 0},
        "started_at": time.time(), "smtp_available": None,
    }
    threading.Thread(target=run_job, args=(job_id, emails), daemon=True).start()
    return jsonify(job_id=job_id, total=len(emails), truncated=truncated)


@app.route("/api/progress/<job_id>")
def progress(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify(error="Unknown job"), 404
    done = job["done"]
    elapsed = time.time() - job["started_at"]
    eta = (elapsed / done) * (job["total"] - done) if done else None
    recent = [r for r in job["results"] if r][-8:] if done else []
    return jsonify(state=job["state"], total=job["total"], done=done,
                   counts=job["counts"], smtp_available=job["smtp_available"],
                   elapsed=round(elapsed, 1), eta=None if eta is None else round(eta),
                   recent=recent)


FIELDS = ["email", "status", "reason", "syntax", "domain", "mx", "catch_all", "smtp"]


@app.route("/api/download/<job_id>")
def download(job_id):
    job = JOBS.get(job_id)
    if not job:
        return jsonify(error="Unknown job"), 404
    only = request.args.get("only")  # valid | invalid | risky | unknown
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=FIELDS, extrasaction="ignore")
    w.writeheader()
    for r in job["results"]:
        if r and (not only or r["status"] == only):
            w.writerow(r)
    name = f"verified_{only or 'all'}_{job_id}.csv"
    return Response(out.getvalue(), mimetype="text/csv",
                    headers={"Content-Disposition": f"attachment; filename={name}"})


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)
